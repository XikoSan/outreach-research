"""Build a separate base-only pilot from verified websites in a prior import.

Dry run by default. Does not call a model, start research, or modify the source run.
"""
import argparse
from copy import copy, deepcopy
from datetime import datetime, timezone
from hashlib import sha256
import json
from pathlib import Path
from openpyxl import load_workbook, Workbook

from app.config import settings
from app.db import Company, Run, Session, SourceRow, log, uid
from app.workbooks import import_rows


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('source_run_id')
    parser.add_argument('offer_run_id')
    parser.add_argument('probe_json', type=Path)
    parser.add_argument('--apply', action='store_true')
    args = parser.parse_args()
    probe = json.loads(args.probe_json.read_text())
    if probe.get('run_id') != args.source_run_id or probe.get('mode') != 'public_fetch_only_no_api':
        raise SystemExit('Probe does not match source run or expected read-only mode')
    recovered = {item['site'] for item in probe.get('www_fallback', []) if item['status'] == 'ok'}
    sites = {item['site'] for item in probe['results'] if item['status'] == 'ok'} | recovered
    if not 1 <= len(sites) <= 30:
        raise SystemExit('Select 1-30 fetchable websites')
    with Session() as db:
        source = db.get(Run, args.source_run_id)
        offer = db.get(Run, args.offer_run_id)
        if not source or not offer or offer.config.get('methodology_version') != 'campaign-v1':
            raise SystemExit('Source run or campaign offer run is missing')
        plan = offer.reports.get('campaign_plan')
        if not plan or not offer.context_snapshot:
            raise SystemExit('Offer run has no saved campaign plan or context')
        companies = db.query(Company).filter_by(run_id=source.id).all()
        selected = [company for company in companies if company.resource in sites]
        if len(selected) != len(sites) or len({company.resource for company in selected}) != len(sites):
            raise SystemExit('Probe sites do not map one-to-one to imported companies')
        selected_ids = {company.id for company in selected}
        rows = db.query(SourceRow).filter_by(run_id=source.id).all()
        coordinates = {(row.sheet, row.row_index) for row in rows if row.company_id in selected_ids}
        print(json.dumps({'sites': len(sites), 'source_rows': len(coordinates), 'apply': args.apply}, ensure_ascii=False))
        if not args.apply:
            return
        workbook = load_workbook(source.source_path)
        subset = Workbook()
        subset.remove(subset.active)
        for sheet in workbook:
            picked = sorted(row_index for title, row_index in coordinates if title == sheet.title)
            if not picked:
                continue
            target = subset.create_sheet(sheet.title)
            for dest_row, source_row in enumerate([1, *picked], 1):
                for cell in sheet[source_row]:
                    new = target.cell(dest_row, cell.column, cell.value)
                    new.data_type = cell.data_type
                    if cell.hyperlink:
                        new._hyperlink = copy(cell.hyperlink)
        workbook.close()
        run_id = uid()
        folder = settings.data_dir / 'uploads' / run_id
        folder.mkdir(parents=True, exist_ok=False)
        path = folder / 'source.xlsx'
        subset.save(path)
        subset.close()
        content = path.read_bytes()
        config = {**offer.config, 'pipeline_version': 3, 'methodology_version': 'campaign-v1',
                  'max_website_pages': 2, 'max_retries': 1, 'concurrency': 1,
                  'research_as_of': datetime.now(timezone.utc).date().isoformat()}
        run = Run(id=run_id, name=f'Пилот базового сбора · {len(sites)} сайтов', filename='base-pilot.xlsx',
                  source_path=str(path.resolve()), source_hash=sha256(content).hexdigest(),
                  sender_snapshot=deepcopy(offer.sender_snapshot), context_snapshot=deepcopy(offer.context_snapshot),
                  reports={'campaign_plan': deepcopy(plan), 'pilot': {'kind': 'base_only',
                          'source_run_id': source.id, 'offer_run_id': offer.id,
                          'probe_file': str(args.probe_json), 'sites': sorted(sites)}},
                  config=config, status='PAUSED', stage='IMPORTED', budget=0)
        db.add(run)
        db.flush()
        import_rows(db, run)
        if run.company_count != len(sites) or run.row_count != len(coordinates):
            db.rollback()
            path.unlink()
            folder.rmdir()
            raise SystemExit('Subset import changed the expected groups or rows')
        log(db, run.id, 'IMPORTED', 'Отдельный пилот базового сбора создан без API-вызовов',
            metrics={'sites': run.company_count, 'rows': run.row_count})
        db.commit()
        print(json.dumps({'run_id': run.id, 'status': run.status,
                          'companies': run.company_count, 'rows': run.row_count}, ensure_ascii=False))


if __name__ == '__main__':
    main()
