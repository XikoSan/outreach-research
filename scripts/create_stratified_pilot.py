"""Create an isolated, reproducible base-only pilot from an existing workbook.

The default is a read-only preview. This command never starts API operations.
The sample has a random core for estimates and a separately labelled set of
hard cases for diagnosis; their results must not be pooled as one random sample.
"""
import argparse
from collections import defaultdict
from copy import copy, deepcopy
from datetime import datetime, timezone
from hashlib import sha256
import json
from pathlib import Path
from random import Random
from urllib.parse import urlsplit

from openpyxl import Workbook, load_workbook
from sqlalchemy import func

from app.config import settings
from app.campaigns import CampaignResearchPlan, VERSION, validate_plan
from app.db import Company, Run, Session, SourceRow, log, uid
from app.workbooks import import_rows, SOCIAL


def select_companies(companies, contact_counts, size, seed):
    """Return (random core, oversampled hard cases) in stable seeded order."""
    rng = Random(seed)
    ordered = sorted(companies, key=lambda company: company.id)
    if size > len(ordered):
        raise ValueError('Размер пилота больше числа компаний в базе')
    core_count = size - min(size // 5, 60)
    core = rng.sample(ordered, core_count)
    used = {company.id for company in core}
    remaining = [company for company in ordered if company.id not in used]
    rng.shuffle(remaining)

    def social_or_missing(company):
        host = (urlsplit(company.resource).hostname or '').lower()
        return company.resource.startswith('unresolved:') or host in SOCIAL

    cohorts = (
        lambda company: social_or_missing(company) or company.group == 'NEEDS_REVIEW',
        lambda company: contact_counts.get(company.id, 0) >= 3,
        lambda company: not (urlsplit(company.resource).hostname or '').lower().endswith('.ru'),
    )
    hard = []
    wanted = size - core_count
    for cohort in cohorts:
        quota = max(1, wanted // len(cohorts))
        for company in remaining:
            if len(hard) >= wanted or quota <= 0:
                break
            if company.id not in used and cohort(company):
                hard.append(company)
                used.add(company.id)
                quota -= 1
    for company in remaining:
        if len(hard) >= wanted:
            break
        if company.id not in used:
            hard.append(company)
            used.add(company.id)
    return core, hard


def frozen_plan_for_source(plan, source):
    if not plan or plan.get('version') != VERSION or plan.get('source_hash') != sha256(source.encode()).hexdigest():
        raise ValueError('Зафиксированный план не соответствует документу с предложением')
    validate_plan(CampaignResearchPlan.model_validate(plan['plan']), source, plan['base_criteria'])
    return deepcopy(plan)


def pilot_config(source_config, offer_config, as_of):
    """Keep the database's acquisition filters, not a former offer pilot's size."""
    return {**offer_config, 'source_filters': source_config.get('source_filters', ''),
            'pipeline_version': 3, 'methodology_version': 'campaign-v1',
            'research_directions': ['website'], 'max_companies': None,
            'max_website_pages': 2, 'max_retries': 1, 'concurrency': 3,
            'research_as_of': as_of}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('source_run_id')
    parser.add_argument('offer_run_id')
    parser.add_argument('--size', type=int, default=200)
    parser.add_argument('--seed', type=int, default=20261012)
    parser.add_argument('--apply', action='store_true')
    parser.add_argument('--frozen-plan-from', help='ID запуска с проверенным планом для точного A/B сравнения')
    args = parser.parse_args()
    if not 20 <= args.size <= 300:
        raise SystemExit('Выберите от 20 до 300 компаний')

    with Session() as db:
        source = db.get(Run, args.source_run_id)
        offer = db.get(Run, args.offer_run_id)
        if not source or not offer or not Path(source.source_path).is_file():
            raise SystemExit('Исходное исследование или контекст не найдены')
        source_text = offer.context_snapshot.get('source_text') or offer.context_snapshot.get('summary', '')
        if not source_text:
            raise SystemExit('В выбранном исследовании нет описания предложения')
        frozen_plan = None
        if args.frozen_plan_from:
            plan_run = db.get(Run, args.frozen_plan_from)
            try:
                frozen_plan = frozen_plan_for_source(
                    plan_run.reports.get('campaign_plan') if plan_run else None, source_text)
            except (ValueError, KeyError) as error:
                raise SystemExit(str(error)) from error
        companies = db.query(Company).filter_by(run_id=source.id).all()
        counts = dict(db.query(SourceRow.company_id, func.count(SourceRow.id))
                      .filter_by(run_id=source.id).group_by(SourceRow.company_id).all())
        try:
            core, hard = select_companies(companies, counts, args.size, args.seed)
        except ValueError as error:
            raise SystemExit(str(error)) from error
        selected = {company.id for company in core + hard}
        coordinates = defaultdict(list)
        origin_by_coordinate = {}
        for sheet, row_index, company_id in db.query(SourceRow.sheet, SourceRow.row_index,
                                                     SourceRow.company_id).filter(
                SourceRow.run_id == source.id, SourceRow.company_id.in_(selected)):
            coordinates[sheet].append(row_index)
            origin_by_coordinate[(sheet, row_index)] = company_id
        row_count = sum(map(len, coordinates.values()))
        preview = {'sampled_companies': len(selected), 'random_core': len(core),
                   'diagnostic_cases': len(hard), 'source_rows': row_count,
                   'seed': args.seed, 'apply': args.apply,
                   'frozen_plan': bool(frozen_plan)}
        print(json.dumps(preview, ensure_ascii=False))
        if not args.apply:
            return

        workbook = load_workbook(source.source_path)
        subset = Workbook()
        subset.remove(subset.active)
        origin_by_pilot_coordinate = {}
        for sheet in workbook:
            picked = sorted(coordinates.get(sheet.title, []))
            if not picked:
                continue
            target = subset.create_sheet(sheet.title)
            for dest_row, source_row in enumerate([1, *picked], 1):
                if dest_row > 1:
                    origin_by_pilot_coordinate[(sheet.title, dest_row)] = origin_by_coordinate[(sheet.title, source_row)]
                for cell in sheet[source_row]:
                    destination = target.cell(dest_row, cell.column, cell.value)
                    destination.data_type = cell.data_type
                    if cell.hyperlink:
                        destination._hyperlink = copy(cell.hyperlink)
        workbook.close()
        if not subset.sheetnames:
            raise SystemExit('Выборка не содержит исходных строк')
        run_id = uid()
        folder = settings.data_dir / 'uploads' / run_id
        folder.mkdir(parents=True, exist_ok=False)
        path = folder / 'source.xlsx'
        subset.save(path)
        subset.close()
        content = path.read_bytes()
        config = pilot_config(source.config, offer.config,
                              datetime.now(timezone.utc).date().isoformat())
        pilot = Run(id=run_id, name=f'Стратифицированный пилот · {args.size} компаний',
                    filename='stratified-pilot.xlsx', source_path=str(path.resolve()),
                    source_hash=sha256(content).hexdigest(),
                    sender_snapshot=deepcopy(offer.sender_snapshot),
                    context_snapshot=deepcopy(offer.context_snapshot),
                    reports={'pilot': {'kind': 'base_only', 'source_run_id': source.id,
                                       'offer_run_id': offer.id, 'seed': args.seed,
                                       'random_core_source_ids': [company.id for company in core],
                                       'diagnostic_source_ids': [company.id for company in hard],
                                       'frozen_plan_run_id': args.frozen_plan_from},
                             **({'campaign_plan': frozen_plan} if frozen_plan else {})},
                    config=config, status='PAUSED', stage='IMPORTED', budget=0)
        try:
            db.add(pilot)
            db.flush()
            import_rows(db, pilot)
            if pilot.row_count != row_count:
                raise ValueError('После импорта изменилась численность исходных строк')
            source_ids_by_pilot_company = defaultdict(set)
            for row in db.query(SourceRow).filter_by(run_id=pilot.id):
                source_ids_by_pilot_company[row.company_id].add(
                    origin_by_pilot_coordinate[(row.sheet, row.row_index)])
            core_ids = {company.id for company in core}
            hard_ids = {company.id for company in hard}
            sample_groups = {'random_core': [], 'diagnostic': [], 'mixed': []}
            for company_id, source_ids in source_ids_by_pilot_company.items():
                label = ('random_core' if source_ids <= core_ids else
                         'diagnostic' if source_ids <= hard_ids else 'mixed')
                sample_groups[label].append(company_id)
            pilot.reports = {**pilot.reports,
                             'pilot': {**pilot.reports['pilot'],
                                       'imported_groups': sample_groups}}
            log(db, pilot.id, 'IMPORTED', 'Стратифицированный пилот создан без API-вызовов',
                metrics={'sampled_companies': args.size, 'imported_companies': pilot.company_count,
                         'source_rows': pilot.row_count})
            db.commit()
        except Exception:
            db.rollback()
            path.unlink(missing_ok=True)
            folder.rmdir()
            raise
        print(json.dumps({'run_id': pilot.id, 'status': pilot.status,
                          'imported_companies': pilot.company_count,
                          'source_rows': pilot.row_count}, ensure_ascii=False))


if __name__ == '__main__':
    main()
