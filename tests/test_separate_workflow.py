import io
import zipfile
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock
import pytest
from openpyxl import load_workbook
from app import credentials
from app.db import Session, Run, Sender, Company, PromptPackage, ModelCall, ResearchContext
from app.llm import Gateway, BudgetPause, UncertainCall
from app.pipeline import Pipeline
from app.schemas import Review
from app.workflow import generate_package
from test_pipeline import FakeGateway, fake_fetch
from test_workflow import workbook


def research(client,monkeypatch,finish=True):
    context=client.post('/api/contexts',json={'name':'Исследования рынков','summary':'Помогаем B2B-компаниям исследовать рынок и проверить позиционирование без гарантии роста.'}).json()
    data=workbook([['Alpha','https://alpha.example','Анна','00123'],['Alpha','https://alpha.example','Анна','00123'],['Beta','https://beta.example','Борис','00777']])
    response=client.post('/api/runs',data={'context_id':context['id'],'name':'Новое исследование','budget':5},files={'file':('contacts.xlsx',data)})
    assert response.status_code==200,response.text
    r=response.json()
    if finish:
        with Session() as db:run=db.get(Run,r['id']);run.status='RUNNING';db.commit()
        monkeypatch.setattr('app.pipeline.fetch',fake_fetch)
        gateway=FakeGateway();Pipeline(r['id'],gateway).execute()
        assert 'prompt' not in gateway.roles
    return r,context


def package(client,rid,**kwargs):
    credentials.write(credentials.GLOBAL_ID,{'key':'sk-offline-test-key'})
    return client.post(f'/api/runs/{rid}/packages',json={'name':'Проверка гипотезы','segment_id':'segment_001','sender':{'brand':'Отправитель','representative':'Иван','signature':'Иван'},'recipient_type':'person','next_step':'Обсудить задачу на встрече','budget':1,**kwargs})


def test_context_reusable_and_frozen_without_sender(client,monkeypatch):
    r,c=research(client,monkeypatch,False)
    assert r['sender_snapshot']=={} and r['context_snapshot']['summary']==c['summary']
    client.put('/api/contexts/'+c['id'],json={'name':'Новое направление','summary':'Другое предложение для следующих исследований.'})
    assert client.get('/api/runs/'+r['id']).json()['context_snapshot']['summary']==c['summary']
    assert client.post('/api/runs',data={'budget':5},files={'file':('missing.xlsx',workbook([['A','example.com','x','y']]))}).status_code==422


def test_upload_preserves_original_and_draft_is_editable(client):
    source='# Предложение\nИсследуем рынки и помогаем проверить продуктовые гипотезы.\nОграничения: не гарантируем продажи.'
    result=client.post('/api/contexts/upload',files={'file':('offer.md',source.encode())})
    assert result.status_code==200
    c=result.json();assert c['source_text']==source and c['summary']
    assert client.get(f"/api/contexts/{c['id']}/source").content==source.encode()
    assert client.put('/api/contexts/'+c['id'],json={'name':'Краткое описание','summary':'Исследование рынка без гарантий продаж.'}).status_code==200
    assert client.get(f"/api/contexts/{c['id']}/source").content==source.encode()
    assert client.post('/api/contexts/upload',files={'file':('bad.docx',b'not a document')}).status_code==422


def test_stage_reports_include_discovery_evidence_and_review(client,monkeypatch):
    r,_=research(client,monkeypatch)
    rows=client.get(f"/api/runs/{r['id']}/companies").json()
    report=client.get(f"/api/runs/{r['id']}/reports?company_id={rows['items'][0]['id']}").json()
    assert report['preparation']['rows']==3 and report['preparation']['companies']==2
    assert report['directions']['website']['discovery']['identity_confirmed']
    assert report['directions']['website']['evidence'][0]['quote']
    assert 'pages' not in report['directions']['website']
    assert report['company']['profile']['signals']
    assert report['segmentation']['reviews'][0]['review']['passed']
    segments=client.get(f"/api/runs/{r['id']}/segments").json()
    assert segments[0]['recommended'] and segments[0]['contact_count']==3


def test_bundle_freezes_selection_sender_and_hypothesis(client,monkeypatch):
    r,_=research(client,monkeypatch)
    contacts=client.get(f"/api/runs/{r['id']}/contacts?segment_id=segment_001").json()
    selected=contacts['items'][0]
    response=package(client,r['id'],selection_mode='contacts',row_ids=[selected['id']],save_sender=True)
    assert response.status_code==200,response.text
    p=response.json();assert p['contact_count']==1
    with Session() as db:
        bundle=db.get(PromptPackage,p['id']);sid=bundle.sender_snapshot['id'];old_hypothesis=bundle.segment_snapshot['hypothesis']
        db.get(Sender,sid).profile={'brand':'Changed'}
        run=db.get(Run,r['id']);run.segments=[]
        bundle.status='RUNNING';db.commit()
    generate_package(p['id'],FakeGateway())
    result=client.get('/api/packages/'+p['id']).json()
    assert result['status']=='COMPLETED' and result['sender']['profile']['brand']=='Отправитель'
    assert result['segment']['hypothesis']==old_hypothesis
    assert client.get('/api/runs/'+r['id']).json()['spent']==0
    contents=client.get('/api/packages/'+p['id']+'/download').content
    with zipfile.ZipFile(io.BytesIO(contents)) as z:
        assert set(z.namelist())=={'contacts.xlsx','prompt.md','snapshot.json'}
        wb=load_workbook(io.BytesIO(z.read('contacts.xlsx')))
        assert list(wb.active.values)[1:]==[('Alpha','https://alpha.example','Анна','00123')]
        assert b'sk-offline' not in z.read('snapshot.json')


def test_all_contacts_keeps_duplicate_rows_and_versions(client,monkeypatch):
    r,_=research(client,monkeypatch)
    first=package(client,r['id']).json();second=package(client,r['id'],recipient_type='company').json()
    assert first['id']!=second['id'] and first['contact_count']==3
    with Session() as db:
        p=db.get(PromptPackage,first['id']);rows=list(load_workbook(Path(p.path)/'contacts.xlsx').active.values)
        assert len(rows)==4 and rows[1]==rows[2]
    assert len(client.get(f"/api/runs/{r['id']}/packages").json())==2


def test_invalid_selection_rejected_and_idempotency_prevents_duplicate_jobs(client,monkeypatch):
    r,_=research(client,monkeypatch)
    assert package(client,r['id'],selection_mode='contacts',row_ids=['not-in-run']).status_code==422
    assert package(client,r['id'],selection_mode='companies',company_ids=[]).status_code==422
    assert package(client,r['id'],recipient_type='unknown').status_code==422
    a=package(client,r['id'],request_id='a'*32);b=package(client,r['id'],request_id='a'*32)
    assert a.status_code==b.status_code==200 and a.json()['id']==b.json()['id']
    assert package(client,r['id'],request_id='a'*32,next_step='Другая задача').status_code==409
    assert len(client.get(f"/api/runs/{r['id']}/packages").json())==1


def test_package_budget_cache_and_unknown_call_reconciliation_are_separate(client,monkeypatch):
    r,_=research(client,monkeypatch);pid=package(client,r['id'],budget=.000001).json()['id']
    provider=Mock();factory=Mock(return_value=provider);monkeypatch.setattr('app.llm.OpenAI',factory)
    with Session() as db:p=db.get(PromptPackage,pid);p.status='RUNNING';db.commit()
    g=Gateway(r['id'],pid)
    provider.responses.create.side_effect=TimeoutError()
    with pytest.raises(UncertainCall):g.ask('synthesis','check',{},Review)
    with Session() as db:
        p=db.get(PromptPackage,pid);assert p.reserved>0;p.status='PAUSED'
        assert db.get(Run,r['id']).reserved==0
        call=db.query(ModelCall).one();assert call.package_id==pid;cid=call.id;db.commit()
    assert client.post('/api/calls/'+cid+'/reconcile',json={'verified_cost':.002,'note':'Сверено вручную'}).status_code==200
    with Session() as db:
        p=db.get(PromptPackage,pid);assert p.reserved==pytest.approx(0);assert p.spent==pytest.approx(.002)
        assert db.get(Run,r['id']).spent==0
        p.status='RUNNING';db.commit()
    provider.responses.create.side_effect=None
    provider.responses.create.return_value=SimpleNamespace(id='ok',usage=SimpleNamespace(input_tokens=10,output_tokens=5),output=[],output_text='{"passed":true,"issues":[]}')
    g.ask('synthesis','check',{},Review);g.ask('synthesis','check',{},Review)
    assert provider.responses.create.call_count==2


def test_prompt_checkpoint_resumes_without_regenerating(client,monkeypatch):
    r,_=research(client,monkeypatch);pid=package(client,r['id']).json()['id']
    with Session() as db:p=db.get(PromptPackage,pid);p.status='RUNNING';db.commit()
    class FailReview(FakeGateway):
        def ask(self,*args,**kwargs):
            if args[0]=='synthesis':raise TimeoutError('review unavailable')
            return super().ask(*args,**kwargs)
    with pytest.raises(TimeoutError):generate_package(pid,FailReview())
    with Session() as db:assert db.get(PromptPackage,pid).prompt
    gateway=FakeGateway();generate_package(pid,gateway)
    assert gateway.roles==['synthesis']
    assert client.get('/api/packages/'+pid).json()['status']=='COMPLETED'


def test_global_key_migration_does_not_restore_deleted_key(client):
    marker=credentials.location(credentials.GLOBAL_ID).parent/'workspace-migrated'
    marker.unlink()
    credentials.write('f'*32,{'key':'old-key'})
    credentials.migrate_global();assert credentials.for_run(None)=='old-key'
    credentials.write(credentials.GLOBAL_ID,{})
    credentials.migrate_global();assert credentials.for_run(None)==''

def test_new_research_without_budget_keeps_import_filters_and_server_models(client):
    from app.api import models
    context=client.post('/api/contexts',json={'name':'Услуги','summary':'Разработка маркетинговой стратегии'}).json()
    content=workbook([['A','example.com','x','y']])
    result=client.post('/api/runs',data={'context_id':context['id'],'name':'База производителей','source_filters':'Москва, производство, 50–200 сотрудников'},files={'file':('base.xlsx',content)})
    assert result.status_code==200
    run=result.json()
    assert run['config']['source_filters']=='Москва, производство, 50–200 сотрудников'
    assert run['config']['models']==models()
    assert run['config']['max_companies'] is None
    assert client.patch('/api/runs/'+run['id']+'/budget',json={'budget':1}).status_code in (404,405)
