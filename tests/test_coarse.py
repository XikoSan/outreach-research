"""End-to-end invariants for quick base preparation, without external API calls."""
from io import BytesIO
import csv
from openpyxl import Workbook,load_workbook
from app.coarse import execute
from app.db import Session,Run,Company,SourceRow,ModelCall


def source_file():
    workbook=Workbook();sheet=workbook.active;sheet.title='Лиды'
    sheet.append(['Название компании','ИНН','Ссылка на сайт','Описание','Email'])
    sheet.append(['Мебель А','1234567890','https://furn-a.example','Производство мебели для офисов','a@example.test'])
    sheet.append(['Мебель А','1234567890','https://furn-a.example','Производство мебели для офисов','b@example.test'])
    sheet.append(['Мебель Б','1234567891','https://furn-b.example','Продажа мебели и диванов','c@example.test'])
    sheet.append(['Без категории','1234567892','https://other.example','Работаем на рынке много лет','d@example.test'])
    sheet.append(['Не писать','1234567893','https://closed.example','Производство мебели','e@example.test'])
    stream=BytesIO();workbook.save(stream);workbook.close();return stream.getvalue()


def test_coarse_import_classification_and_export(client):
    content=source_file()
    response=client.post('/api/coarse-runs',files={'file':('leads.xlsx',content,'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')},data={'name':'Тест базы'})
    assert response.status_code==200,response.text
    run_id=response.json()['id']
    assert response.json()['row_count']==5 and response.json()['company_count']==4
    assert response.json()['context_snapshot']=={}
    assert client.post(f'/api/runs/{run_id}/start').status_code==200
    with Session() as db:
        blocked=db.query(Company).filter_by(run_id=run_id,name='Не писать').one()
        blocked.group='EXCLUDED';blocked.profile={'exclusion':{'do_not_send':True,'reason':'Подтверждено вручную'}}
        db.commit()
    execute(run_id)
    result=client.get(f'/api/runs/{run_id}').json()
    assert result['status']=='COMPLETED' and result['spent']==0 and result['reserved']==0
    assert result['finished']==4 and result['groups']['segment_001']==2
    assert result['groups']['GENERAL']==1 and result['groups']['DO_NOT_SEND']==1
    assert client.get(f'/api/coarse-runs/{run_id}/companies?group=DO_NOT_SEND').json()['total']==1
    assert client.get(f'/api/runs/{run_id}/companies?group=segment_001').status_code==200
    export=client.post(f'/api/runs/{run_id}/exports')
    assert export.status_code==200,export.text
    manifest=export.json()['manifest'];eid=export.json()['id']
    assert manifest['rows']==5 and manifest['mailing_rows']==4 and manifest['do_not_send_rows']==1
    assert manifest['groups']=={'DO_NOT_SEND':1,'GENERAL':1,'segment_001':3}
    data=client.get(f'/api/exports/{eid}/download?file=integration.csv').content.decode('utf-8-sig').splitlines()
    indexed=list(csv.DictReader(data))
    assert len(indexed)==5 and len({r['source_row_id'] for r in indexed})==5
    assert len({r['company_id'] for r in indexed if r['audience_id']=='segment_001'})==2
    assert sum(int(r['do_not_send']) for r in indexed)==1
    mailing=client.get(f'/api/exports/{eid}/download?file=audiences/segment_001.xlsx').content
    original=load_workbook(BytesIO(mailing))
    assert [original.active.cell(i,5).value for i in range(2,5)]==['a@example.test','b@example.test','c@example.test']
    original.close()
    single=client.get(f'/api/exports/{eid}/download?file=mailing.xlsx').content
    ready=load_workbook(BytesIO(single))
    assert ready.active.max_row==5  # header + four rows eligible for mailing
    assert ready.active.cell(1,11).value=='audience_id'
    assert [ready.active.cell(i,11).value for i in range(2,6)].count('segment_001')==3
    assert not any(ready.active.cell(i,5).value=='e@example.test' for i in range(2,6))
    ready.close()
    with Session() as db:
        changed=db.query(Company).filter_by(run_id=run_id,name='Мебель Б').one().id
    review=client.post(f'/api/coarse-runs/{run_id}/companies/{changed}/review',json={'audience_id':'GENERAL','reason':'Ручная проверка показала другую аудиторию'})
    assert review.status_code==200 and review.json()['segment_set_version']==2
    assert client.get(f'/api/exports/{eid}/download?file=mailing.xlsx').status_code==409
    refreshed=client.post(f'/api/runs/{run_id}/exports').json()['manifest']
    assert refreshed['groups']['segment_001']==2 and refreshed['groups']['GENERAL']==2
    with Session() as db:
        assert db.query(ModelCall).filter_by(run_id=run_id).count()==0
        assert db.query(SourceRow).filter_by(run_id=run_id).count()==5
        assert db.get(Run,run_id).budget==0
