"""First pass keeps the offer grounded and a missing vacancy neutral."""
from io import BytesIO
from openpyxl import Workbook,load_workbook

from app import offer_coarse
from app.db import Session,Run,Company,ModelCall


def workbook():
    wb=Workbook();ws=wb.active;ws.title='Лиды'
    ws.append(['Название компании','ИНН','Ссылка на сайт','Описание','Email'])
    ws.append(['Мебель А','1234567890','https://a.example','Производство мебели для офисов','a@example.test'])
    ws.append(['Мебель Б','1234567891','https://b.example','Производство мебели и столов','b@example.test'])
    ws.append(['Неясная','1234567892','https://c.example','Работаем давно','c@example.test'])
    stream=BytesIO();wb.save(stream);wb.close();return stream.getvalue()


def test_offer_first_pass_and_vacancy_are_separate(client,monkeypatch):
    offer='Развиваем продажи производителей мебели и ищем новые каналы сбыта.'
    response=client.post('/api/offer-runs',files={'file':('base.xlsx',workbook(),'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')},data={'name':'Оффер и база','offer_text':offer})
    assert response.status_code==200,response.text
    rid=response.json()['id']
    listing=client.get('/api/offer-runs').json()
    assert listing['total']==1 and listing['items'][0]['id']==rid
    assert response.json()['context_snapshot']['source_text']==offer
    plan=offer_coarse.RoutingPlan(audiences=[offer_coarse.AudiencePlan(name='Производители мебели',offer_quote=offer,
        activity_stems=['мебел'],vacancy_roles=['менеджер по продажам'],first_letter_angle='Предложить обсудить новые каналы сбыта')])
    monkeypatch.setattr(offer_coarse.Gateway,'ask',lambda *args,**kwargs:plan)
    monkeypatch.setattr(offer_coarse,'vacancy_signal',lambda url,roles:{'state':'CONFIRMED','role':'Менеджер по продажам','url':'https://a.example/jobs/sales','retrieved_at':'2026-10-16'} if 'a.example' in url else {'state':'NOT_CONFIRMED'})
    with Session() as db:
        run=db.get(Run,rid);run.status='RUNNING';db.commit()
    offer_coarse.execute(rid)
    result=client.get(f'/api/runs/{rid}').json()
    assert result['status']=='COMPLETED' and result['groups']=={'segment_001':2,'GENERAL':1}
    assert result['confirmed_vacancies']==1
    assert result['segments'][0]['offer_quote']==offer
    assert result['spent']==0  # the model was replaced by a deterministic test stub
    companies=client.get(f'/api/coarse-runs/{rid}/companies').json()['items']
    assert {c['name']:c['audience_id'] for c in companies}=={'Мебель А':'segment_001','Мебель Б':'segment_001','Неясная':'GENERAL'}
    assert {c['name']:c['vacancy']['state'] for c in companies}=={'Мебель А':'CONFIRMED','Мебель Б':'NOT_CONFIRMED','Неясная':'NOT_CONFIRMED'}
    exported=client.post(f'/api/runs/{rid}/exports')
    assert exported.status_code==200,exported.text
    eid=exported.json()['id']
    ready=load_workbook(BytesIO(client.get(f'/api/exports/{eid}/download?file=mailing.xlsx').content))
    headers=[c.value for c in ready.active[1]]
    assert headers[-3:]==['vacancy_signal','vacancy_role','vacancy_url']
    rows=list(ready.active.values)
    assert [row[-3] for row in rows[1:]].count('CONFIRMED')==1
    ready.close()
    with Session() as db:
        assert db.query(ModelCall).filter_by(run_id=rid).count()==0


def test_vacancy_requires_company_site_link_and_role(monkeypatch):
    def fake_fetch(url):
        return {'url':url,'title':'Карьера' if '/career' in url else 'Главная','retrieved_at':'2026-10-16',
                'links':([{'url':'https://a.example/career','label':'Карьера'}] if url=='https://a.example' else
                         [{'url':'https://a.example/jobs/sales','label':'Менеджер по продажам'},
                          {'url':'https://external.example/jobs/sales','label':'Менеджер по продажам'}])}
    monkeypatch.setattr(offer_coarse,'fetch',fake_fetch)
    assert offer_coarse.vacancy_signal('https://a.example',['менеджер по продажам'])['url']=='https://a.example/jobs/sales'
    assert offer_coarse.vacancy_signal('https://a.example',['маркетолог'])['state']=='NOT_CONFIRMED'
    assert offer_coarse.vacancy_signal('unresolved:abc',['менеджер по продажам'])['state']=='NOT_CONFIRMED'


def test_plan_cannot_invent_offer_fact():
    plan=offer_coarse.RoutingPlan(audiences=[offer_coarse.AudiencePlan(name='Группа',offer_quote='Несуществующая услуга',activity_stems=['мебел'],vacancy_roles=[],first_letter_angle='Предложение сотрудничества')])
    try:offer_coarse.validate_plan(plan,'Реальная услуга по продажам мебели.', ['мебели'])
    except ValueError as error:assert 'цитатой' in str(error)
    else:assert False,'The invented offer must be rejected'


def test_offer_is_required(client):
    response=client.post('/api/offer-runs',files={'file':('base.xlsx',workbook(),'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')},data={'name':'Без оффера','offer_text':'нет'})
    assert response.status_code==422
    assert client.get('/api/offer-runs').json()['total']==0
