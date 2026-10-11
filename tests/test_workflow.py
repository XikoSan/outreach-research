import io,zipfile
from openpyxl import Workbook,load_workbook
from app.db import Session,Run,Company,SourceRow,Export
from app.rules import classify
from app.workbooks import canonical_resource
from app.crawler import public_address,SourceUnavailable
from app.pipeline import verify_report
from app.schemas import ResearchReport
import pytest

def workbook(rows,headers=None):
    w=Workbook();s=w.active;s.title='Сотрудники';s.append(headers or ['Название компании','Ссылка на сайт','Имя','Телефон'])
    for r in rows:s.append(r)
    s['D2'].number_format='@'
    b=io.BytesIO();w.save(b);return b.getvalue()

def upload(client,content):
    sid=client.get('/api/senders').json()[0]['id']
    return client.post('/api/runs',data={'sender_id':sid,'budget':5},files={'file':('input.xlsx',content)})

def test_roundtrip_preserves_duplicates_and_strings(client):
    rows=[['Альфа','https://example.com','Анна','00123'],['Альфа','http://www.example.com/path','Борис','00123'],['Без сайта',None,'Вера','00456'],['Альфа','https://example.com','Анна','00123']]
    r=upload(client,workbook(rows));assert r.status_code==200,r.text
    run=r.json();assert run['row_count']==4 and run['company_count']==2
    res=client.post(f"/api/runs/{run['id']}/exports");assert res.status_code==200,res.text
    ex=res.json();assert ex['manifest']['rows']==4
    z=zipfile.ZipFile(io.BytesIO(client.get(f"/api/exports/{ex['id']}/download").content))
    actual=[]
    for f in z.namelist():
        w=load_workbook(io.BytesIO(z.read(f)));actual+=list(w.active.values)[1:]
    assert sorted(map(str,actual))==sorted(map(str,map(tuple,rows)))
    assert not any(n.endswith('.md') for n in z.namelist())
    assert client.get(f"/api/exports/{ex['id']}/download?file=../../.env").status_code==404

def test_formula_rejected_before_research(client):
    r=upload(client,workbook([['A','https://example.com','A','=1+1']]))
    assert r.status_code==422
    assert client.get('/api/runs').json()==[]

def test_duplicate_headers_preserved(client):
    r=upload(client,workbook([['A','https://example.com','x','y']],['Название компании','Ссылка на сайт','X','X']));assert r.status_code==200
    eid=client.post('/api/runs/'+r.json()['id']+'/exports').json()['id']
    z=zipfile.ZipFile(io.BytesIO(client.get('/api/exports/'+eid+'/download').content));w=load_workbook(io.BytesIO(z.read(z.namelist()[0])))
    assert list(w.active.values)==[('Название компании','Ссылка на сайт','X','X'),('A','https://example.com','x','y')]

def test_sender_snapshot_is_immutable(client):
    r=upload(client,workbook([['A','example.com','x','y']])).json()
    p=client.get('/api/senders').json()[0]['profile'];p['representative']='Другой отправитель'
    assert client.post('/api/senders',json=p).status_code==200
    assert client.get('/api/runs/'+r['id']).json()['sender_snapshot']['profile']['representative']=='Артём Кучеренко'

def test_start_without_key_is_explicit(client):
    from app.config import settings
    previous=settings.openai_api_key;settings.openai_api_key=''
    try:
        r=upload(client,workbook([['A','example.com','x','y']])).json()
        assert client.post('/api/runs/'+r['id']+'/start').status_code==409
    finally:settings.openai_api_key=previous

def test_social_pages_not_merged_by_host():
    assert canonical_resource('https://vk.com/alpha')!=canonical_resource('https://vk.com/beta')
    assert canonical_resource('https://vk.com') is None
    assert canonical_resource('https://user:password@example.com') is None
    assert canonical_resource('https://linkedin.com/in/person',True) is None
    assert canonical_resource('https://youtube.com/watch?v=123',True) is None
    assert canonical_resource('https://x.com/brand/status/123',True) is None
    assert canonical_resource('https://youtu.be/abc123',True) is None

def test_corporate_social_fallback_uses_all_supported_columns(client):
    headers=['Название компании','Ссылка на сайт','Telegram компании','Instagram компании','LinkedIn сотрудника компании']
    rows=[
        ['Альфа',None,'https://t.me/alpha','https://instagram.com/alpha','https://linkedin.com/in/person'],
        ['Альфа',None,'https://t.me/alpha',None,None],
        ['Бета',None,None,'https://instagram.com/beta',None],
        ['Гамма',None,None,None,'https://linkedin.com/in/person'],
    ]
    response=upload(client,workbook(rows,headers));assert response.status_code==200,response.text
    run=response.json();assert run['row_count']==4 and run['company_count']==3
    with Session() as db:
        companies={c.name:c for c in db.query(Company).filter_by(run_id=run['id'])}
        assert companies['Альфа'].resource=='https://t.me/alpha'
        assert companies['Альфа'].context['socials']==['https://t.me/alpha','https://instagram.com/alpha']
        assert companies['Бета'].resource=='https://instagram.com/beta'
        assert companies['Гамма'].group=='NEEDS_REVIEW'


def test_official_page_and_x_company_headers_are_supported(client):
    headers=['Название компании','Ссылка на сайт','Официальная страница компании','X компании']
    rows=[['Дельта',None,'https://www.tiktok.com/@delta','https://x.com/delta']]
    response=upload(client,workbook(rows,headers));assert response.status_code==200,response.text
    with Session() as db:
        company=db.query(Company).filter_by(run_id=response.json()['id']).one()
        assert company.resource=='https://tiktok.com/@delta'
        assert company.context['socials']==['https://tiktok.com/@delta','https://x.com/delta']


def test_website_stays_anchor_and_later_social_links_are_preserved(client):
    headers=['Название компании','Ссылка на сайт','Соцсети компании','Имя']
    rows=[
        ['Альфа','https://alpha.example/about','https://t.me/alpha','Первый'],
        ['Альфа','http://www.alpha.example/products','https://youtube.com/@alpha https://instagram.com/alpha','Второй'],
    ]
    response=upload(client,workbook(rows,headers));assert response.status_code==200,response.text
    run=response.json();assert run['row_count']==2 and run['company_count']==1
    with Session() as db:
        company=db.query(Company).filter_by(run_id=run['id']).one()
        assert company.resource=='https://alpha.example'
        assert company.context['socials']==['https://t.me/alpha','https://youtube.com/@alpha','https://instagram.com/alpha']
        assert db.query(SourceRow).filter_by(run_id=run['id'],company_id=company.id).count()==2


def fixtures():
    signals=[{'name':'products','state':'confirmed','confidence':'HIGH','evidence_ids':['a']},{'name':'audiences','state':'confirmed','confidence':'HIGH','evidence_ids':['b']}]
    evidence=[{'id':'a','url':'https://a.com/x','quote':'several product lines','scope':'company'},{'id':'b','url':'https://b.com/x','quote':'several target groups','scope':'product'}]
    segment={'id':'segment_001','required_signals':['products'],'supporting_signals':['audiences'],'contradicting_signals':['separate_messages']}
    return {'signals':signals},evidence,segment

def test_classification_requires_independent_origins():
    p,e,s=fixtures();assert classify(p,e,[s])[0]=='segment_001'
    e[1]['url']='https://a.com/other';assert classify(p,e,[s])[0]=='UNCLASSIFIED'

def test_unknown_low_and_contradiction_never_force_segment():
    p,e,s=fixtures();p['signals'][0]['state']='unknown';assert classify(p,e,[s])[0]=='INSUFFICIENT_DATA'
    p,e,s=fixtures();p['signals'].append({'name':'separate_messages','state':'confirmed','confidence':'HIGH','evidence_ids':['b']});assert classify(p,e,[s])[0]=='UNCLASSIFIED'
    p,e,s=fixtures();assert classify(p,e,[s,{**s,'id':'segment_002'}])[0]=='AMBIGUOUS'

def test_category_is_not_company_evidence():
    p,e,s=fixtures();e[0]['scope']='category';assert classify(p,e,[s])[0]!='segment_001'

def test_non_public_urls_blocked():
    for url in ['http://127.0.0.1','http://169.254.169.254','file:///etc/passwd','http://localhost']:
        with pytest.raises(SourceUnavailable):public_address(url)

def test_unsupported_quote_is_dropped():
    report=ResearchReport(summary='test',evidence=[dict(id='1',url='https://example.com',quote='this quotation was invented',published_at=None,scope='company',origin='example.com')],signals=[dict(name='x',state='confirmed',confidence='HIGH',evidence_ids=['1'],explanation='')],contradictions=[],limitations=[])
    result=verify_report(report,[{'url':'https://example.com','text':'something else entirely'}],'website',30)
    assert result['evidence']==[] and result['signals']==[]

def test_old_social_cannot_be_primary():
    report=ResearchReport(summary='test',evidence=[dict(id='1',url='https://example.com',quote='there are multiple audiences',published_at='2020-01-01',scope='company',origin='example.com')],signals=[dict(name='x',state='confirmed',confidence='HIGH',evidence_ids=['1'],explanation='')],contradictions=[],limitations=[])
    result=verify_report(report,[{'url':'https://example.com','text':'there are multiple audiences'}],'social',30)
    assert result['signals'][0]['confidence']=='LOW'

def test_cross_origin_write_rejected(client):
    assert client.post('/api/senders',headers={'Origin':'https://untrusted.example'},json={}).status_code==403
