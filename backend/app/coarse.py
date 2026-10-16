"""Fast, auditable audience preparation from the uploaded workbook only.

This deliberately classifies business activity, never a company's presumed problem.
It does not call an LLM or crawl websites. Later enrichment can be measured separately.
"""
from __future__ import annotations

from collections import Counter, defaultdict
from copy import copy
import csv
import hashlib
import json
from pathlib import Path
import re
import shutil
import zipfile

from openpyxl import Workbook, load_workbook

from .db import Company, Export, Run, Session, SourceRow, log, now, uid
from .workbooks import fingerprint

# Terms describe broad activity. Ambiguous or overlapping descriptions stay GENERAL.
AUDIENCES = (
    ('segment_001', 'Мебель и интерьер', (r'мебел', r'диван', r'кухн(?:и|я|ый)', r'матрас', r'шкаф', r'интерьер')),
    ('segment_002', 'Строительные материалы', (r'кирпич', r'бетон', r'цемент', r'кровл', r'фасад', r'стройматериал', r'пиломатериал', r'утеплител')),
    ('segment_003', 'Металлоизделия и обработка', (r'металлоиздел', r'металлообработ', r'металлоконструкц', r'листов(?:ой|ого) металл', r'нержавеющ', r'сварочн')),
    ('segment_004', 'Промышленное оборудование', (r'станк', r'насос', r'компрессор', r'конвейер', r'промышленн.{0,20}оборудован', r'производственн.{0,20}оборудован')),
    ('segment_005', 'Электротехника и автоматика', (r'электротех', r'электрооборудован', r'кабел', r'светильник', r'автоматизац', r'электромонтаж')),
    ('segment_006', 'Пищевые продукты и производство', (r'пищев', r'продукт(?:ы|ов) питания', r'кондитер', r'хлебобулоч', r'молочн', r'мясопереработ', r'напитк')),
    ('segment_007', 'Одежда и текстиль', (r'одежд', r'текстил', r'ткан', r'обув', r'трикотаж', r'спецодежд')),
    ('segment_008', 'IT и цифровые продукты', (r'разработк.{0,20}программ', r'программное обеспечен', r'it.компан', r'веб.разработ', r'автоматизац.{0,20}бизнес', r'информационн.{0,20}систем')),
    ('segment_009', 'Логистика и транспорт', (r'логистик', r'грузоперевоз', r'транспортн.{0,20}услуг', r'складск.{0,20}услуг', r'экспедирован')),
)


def classify(company: Company) -> dict:
    """Return one conservative broad audience and the basis of the assignment."""
    ctx=company.context or {}
    if company.group=='EXCLUDED' or (company.profile or {}).get('exclusion',{}).get('do_not_send'):
        return {'audience_id':'DO_NOT_SEND','confidence':'HIGH','reason':'Подтверждённое исключение','source':'reviewed_exclusion','matched_terms':[]}
    if ctx.get('identity_conflict'):
        return {'audience_id':'GENERAL','confidence':'LOW','reason':'Неоднозначная связь контактов с компанией','source':'database','matched_terms':[]}
    description=str(ctx.get('description') or '').lower().replace('ё','е')
    name=str(company.name or '').lower().replace('ё','е')
    if not description.strip() and not name.strip():
        return {'audience_id':'GENERAL','confidence':'LOW','reason':'Нет описания деятельности','source':'database','matched_terms':[]}
    # A description is more reliable for business activity than a brand name.
    text=description or name
    matches=[]
    for audience_id,_,patterns in AUDIENCES:
        terms=[p for p in patterns if re.search(p,text,re.I)]
        if terms:matches.append((audience_id,terms))
    if len(matches)!=1:
        return {'audience_id':'GENERAL','confidence':'LOW','reason':'Нет однозначной широкой категории' if matches else 'Нет признаков деятельности для широкой категории','source':'database_description' if description else 'database_name','matched_terms':[]}
    audience_id,terms=matches[0]
    return {'audience_id':audience_id,'confidence':'MEDIUM' if description else 'LOW','reason':'Явная категория деятельности в '+('описании базы' if description else 'названии'),
            'source':'database_description' if description else 'database_name','matched_terms':terms}


def execute(run_id: str) -> None:
    """Idempotent per company; a pause leaves completed classifications intact."""
    with Session() as db:
        run=db.get(Run,run_id)
        if not run or run.config.get('workflow')!='coarse-v1':raise ValueError('Неизвестная база быстрого анализа')
        run.stage='COARSE_CLASSIFICATION';db.commit()
    with Session() as db:
        ids=[x[0] for x in db.query(Company.id).filter_by(run_id=run_id).order_by(Company.id)]
    for i,company_id in enumerate(ids):
        with Session() as db:
            run=db.get(Run,run_id)
            if run.status in ('PAUSING','PAUSED'):
                run.status='PAUSED';db.commit();return
            company=db.get(Company,company_id)
            if company.stage=='COARSE_READY':continue
            result=classify(company)
            company.profile={**(company.profile or {}),'coarse':result}
            company.group=result['audience_id']
            company.reason=result['reason']
            company.identity='NEEDS_REVIEW' if company.context.get('identity_conflict') else 'BASE_LINKED'
            company.stage='COARSE_READY'
            if i%100==0:log(db,run_id,'COARSE','Классификация базы',metrics={'processed':i+1,'total':len(ids)})
            db.commit()
    with Session() as db:
        run=db.get(Run,run_id)
        if run.status in ('PAUSING','PAUSED'):
            run.status='PAUSED';db.commit();return
        companies=db.query(Company).filter_by(run_id=run_id).all()
        tally=Counter(c.group for c in companies)
        minimum=max(2,(len(companies)+99)//100)
        valid={audience_id for audience_id,_,_ in AUDIENCES if tally[audience_id]>=minimum}
        for company in companies:
            if company.group.startswith('segment_') and company.group not in valid:
                company.group='GENERAL'
                company.reason='Категория слишком мала для отдельной массовой аудитории'
                company.profile={**company.profile,'coarse':{**company.profile['coarse'],'audience_id':'GENERAL','reason':company.reason}}
        tally=Counter(c.group for c in companies)
        run.segments=[{'id':aid,'name':label,'definition':'Компания с явным указанием этого направления деятельности в исходной базе','hypothesis':'Для этой аудитории допустимо общее обращение по сфере деятельности без утверждений о проблемах и потребностях','source':'database','company_count':tally[aid]}
                      for aid,label,_ in AUDIENCES if aid in valid]
        run.segmentation_version=1
        run.reports={**(run.reports or {}),'coarse':{'version':'coarse-v1','processed_at':now(),'minimum_audience_size':minimum,'groups':dict(tally),'method':'database_rules','api_cost':run.spent}}
        run.stage='COARSE_READY';run.status='COMPLETED';run.error=None
        log(db,run_id,'COARSE','Быстрый анализ завершён',metrics={'companies':len(companies),'groups':dict(tally),'api_cost':run.spent})
        db.commit()


def export(db, run: Run, root: Path) -> Export:
    """Export original rows and stable IDs; excluded rows are never in mailing files."""
    if run.config.get('workflow') not in ('coarse-v1','offer-coarse-v2') or run.status!='COMPLETED':raise ValueError('Быстрый анализ не завершён')
    if hashlib.sha256(Path(run.source_path).read_bytes()).hexdigest()!=run.source_hash:raise ValueError('Исходный файл изменён')
    companies={c.id:c for c in db.query(Company).filter_by(run_id=run.id)}
    rows=db.query(SourceRow).filter_by(run_id=run.id).order_by(SourceRow.sheet,SourceRow.row_index).all()
    if len(rows)!=run.row_count or len(companies)!=run.company_count:raise ValueError('Число исходных строк изменилось')
    grouped=defaultdict(list)
    for row in rows:
        group=companies[row.company_id].group
        if group not in {'GENERAL','DO_NOT_SEND'} and group not in {s['id'] for s in run.segments}:raise ValueError('Неизвестная аудитория')
        grouped[group].append(row)
    export_id=uid();stage=Path(root)/('tmp-'+export_id);dest=Path(root)/export_id
    stage.mkdir(parents=True)
    source=load_workbook(run.source_path)
    manifest={'kind':run.config.get('workflow','coarse-v1'),'contract_version':1,'source_run_id':run.id,'export_id':export_id,'segment_set_version':run.segmentation_version,'rows':len(rows),'files':[],'groups':{},'mailing_rows':0,'do_not_send_rows':len(grouped['DO_NOT_SEND'])}
    seen=set()
    try:
        for group,members in sorted(grouped.items()):
            out=Workbook();out.remove(out.active)
            by_sheet=defaultdict(list)
            for row in members:by_sheet[row.sheet].append(row)
            for spec in run.config['sheets']:
                title=spec['name'];src=source[title];ws=out.create_sheet(title)
                ws.freeze_panes='A2'
                for key,dim in src.column_dimensions.items():ws.column_dimensions[key]=copy(dim)
                original_indices=[1]+[row.row_index for row in by_sheet[title]]
                for dest_i,original_i in enumerate(original_indices,1):
                    for cell in src[original_i]:
                        target=ws.cell(dest_i,cell.column,cell.value);target.data_type=cell.data_type
                        if cell.has_style:target._style=copy(cell._style)
                        if cell.hyperlink:target._hyperlink=copy(cell.hyperlink)
                        if cell.comment:target.comment=copy(cell.comment)
                    ws.row_dimensions[dest_i].height=src.row_dimensions[original_i].height
                ws.auto_filter.ref=ws.dimensions
            folder='excluded' if group=='DO_NOT_SEND' else 'audiences'
            path=stage/folder/(group+'.xlsx');path.parent.mkdir(exist_ok=True);out.save(path);out.close()
            check=load_workbook(path)
            for spec in run.config['sheets']:
                title=spec['name'];ws=check[title];src=source[title]
                if fingerprint(ws[1])!=fingerprint(src[1]):raise ValueError('Заголовок изменился')
                for idx,row in enumerate(by_sheet[title],2):
                    if fingerprint(ws[idx])!=row.fingerprint:raise ValueError(f'Строка {row.id} изменилась')
                    if row.id in seen:raise ValueError('Дублирование исходной строки')
                    seen.add(row.id)
            check.close();manifest['files'].append(str(path.relative_to(stage)));manifest['groups'][group]=len(members)
            if group!='DO_NOT_SEND':manifest['mailing_rows']+=len(members)
        if len(seen)!=run.row_count:raise ValueError('Потеряна исходная строка')
        # One import-ready workbook for an external sender. Original cells remain in order;
        # immutable IDs travel in the same row, so replies can be linked without name guessing.
        mailing=Workbook();mailing.remove(mailing.active)
        service_headers=['source_run_id','export_id','segment_set_version','company_id','source_row_id','audience_id']
        if run.config.get('workflow')=='offer-coarse-v2':service_headers+=['vacancy_signal','vacancy_role','vacancy_url']
        for spec in run.config['sheets']:
            title=spec['name'];original=source[title];sheet=mailing.create_sheet(title)
            width=original.max_column
            for cell in original[1]:
                target=sheet.cell(1,cell.column,cell.value)
                if cell.has_style:target._style=copy(cell._style)
            for column,label in enumerate(service_headers,width+1):sheet.cell(1,column,label)
            source_rows=[row for row in rows if row.sheet==title and companies[row.company_id].group!='DO_NOT_SEND']
            for index,row in enumerate(source_rows,2):
                for cell in original[row.row_index]:
                    target=sheet.cell(index,cell.column,cell.value);target.data_type=cell.data_type
                    if cell.has_style:target._style=copy(cell._style)
                    if cell.hyperlink:target._hyperlink=copy(cell.hyperlink)
                vacancy=(companies[row.company_id].profile or {}).get('coarse_vacancy',{})
                values=[run.id,export_id,run.segmentation_version,row.company_id,row.id,companies[row.company_id].group]
                if run.config.get('workflow')=='offer-coarse-v2':values += [vacancy.get('state','NOT_CONFIRMED'),vacancy.get('role',''),vacancy.get('url','')]
                for column,value in enumerate(values,width+1):
                    sheet.cell(index,column,value)
            sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
        mailing_path=stage/'mailing.xlsx';mailing.save(mailing_path);mailing.close()
        manifest['files'].append('mailing.xlsx')
        path=stage/'integration.csv'
        with path.open('w',newline='',encoding='utf-8-sig') as stream:
            writer=csv.writer(stream)
            headers=['source_run_id','export_id','segment_set_version','company_id','source_row_id','audience_id','do_not_send','source_sheet','source_row_index']
            if run.config.get('workflow')=='offer-coarse-v2':headers+=['vacancy_signal','vacancy_role','vacancy_url']
            writer.writerow(headers)
            for row in rows:
                company=companies[row.company_id]
                values=[run.id,export_id,run.segmentation_version,company.id,row.id,company.group,int(company.group=='DO_NOT_SEND'),row.sheet,row.row_index]
                if run.config.get('workflow')=='offer-coarse-v2':
                    vacancy=(company.profile or {}).get('coarse_vacancy',{})
                    values += [vacancy.get('state','NOT_CONFIRMED'),vacancy.get('role',''),vacancy.get('url','')]
                writer.writerow(values)
        manifest['files'].append('integration.csv')
        description=stage/'audiences.json'
        description.write_text(json.dumps(run.segments,ensure_ascii=False,indent=2),encoding='utf-8')
        manifest['files'].append('audiences.json')
        manifest['checksums']={f:hashlib.sha256((stage/f).read_bytes()).hexdigest() for f in manifest['files']}
        with zipfile.ZipFile(stage/'results.zip','w',zipfile.ZIP_DEFLATED) as archive:
            for filename in manifest['files']:archive.write(stage/filename,filename)
        source.close();stage.rename(dest)
        record=Export(id=export_id,run_id=run.id,path=str(dest),manifest=manifest)
        db.add(record);db.commit();return record
    except Exception:
        source.close();shutil.rmtree(stage,ignore_errors=True);raise
