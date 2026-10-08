"""Research context, stage reports, audience selection and independent prompt jobs."""
import hashlib
import io
import json
import re
import shutil
import zipfile
from copy import copy
from pathlib import Path
from typing import Literal
from xml.etree import ElementTree
from fastapi import APIRouter, HTTPException, UploadFile, File, Query
from fastapi.responses import FileResponse
from pydantic import BaseModel, Field
from openpyxl import load_workbook, Workbook
from sqlalchemy import func
from . import credentials
from .config import settings
from .db import Session, Run, Company, SourceRow, Sender, ResearchContext, PromptPackage, ModelCall, uid, now, log
from .schemas import SenderProfile, PromptResult, Review
from .llm import Gateway, RunPaused
from .prompts import PROMPT, REVIEW
from .workbooks import fingerprint
from .presentation import company_view

router=APIRouter(prefix='/api')
ACTIVE=('RUNNING','QUEUED','PAUSING')

def get(db,model,id):
    value=db.get(model,id)
    if not value:raise HTTPException(404,'Не найдено')
    return value

def context_view(c):
    return {k:getattr(c,k) for k in ('id','name','summary','source_name','source_text','created_at')}

class ContextInput(BaseModel):
    name:str=Field(min_length=1,max_length=200)
    summary:str=Field(min_length=10,max_length=8000)

@router.get('/contexts')
def contexts():
    with Session() as db:return [context_view(c) for c in db.query(ResearchContext).order_by(ResearchContext.created_at.desc())]

@router.post('/contexts')
def create_context(payload:ContextInput):
    if not payload.name.strip() or len(payload.summary.strip())<10:raise HTTPException(422,'Заполните название и описание предложения')
    with Session() as db:
        c=ResearchContext(name=payload.name.strip(),summary=payload.summary.strip());db.add(c);db.commit();return context_view(c)

@router.put('/contexts/{cid}')
def edit_context(cid:str,payload:ContextInput):
    if not payload.name.strip() or len(payload.summary.strip())<10:raise HTTPException(422,'Заполните название и описание предложения')
    with Session() as db:
        c=get(db,ResearchContext,cid);c.name=payload.name.strip();c.summary=payload.summary.strip();db.commit();return context_view(c)

@router.post('/contexts/upload')
def upload_context(file:UploadFile=File(...)):
    content=file.file.read(5_000_001)
    if len(content)>5_000_000:raise HTTPException(413,'Лимит документа 5 МБ')
    suffix=Path(file.filename or '').suffix.lower()
    try:
        if suffix in ('.txt','.md'):text=content.decode('utf-8-sig')
        elif suffix=='.docx':
            with zipfile.ZipFile(io.BytesIO(content)) as z:
                if z.getinfo('word/document.xml').file_size>10_000_000:raise ValueError()
                root=ElementTree.fromstring(z.read('word/document.xml'))
                ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
                text='\n'.join(''.join(t.text or '' for t in p.findall('.//w:t',ns)) for p in root.findall('.//w:p',ns))
        else:raise ValueError()
        if not 10<=len(text.strip())<=100_000:raise ValueError()
    except Exception:raise HTTPException(422,'Нужен TXT, MD или DOCX с текстом от 10 до 100 000 символов')
    # An extractive draft is free and retains wording. It must be reviewed before use.
    lines=[line.strip() for line in text.splitlines() if line.strip()]
    draft=[];size=0
    for line in lines:
        if size+len(line)>2500:break
        draft.append(line);size+=len(line)+1
    summary='\n'.join(draft) or text[:2500]
    cid=uid();folder=settings.data_dir/'materials'/cid;folder.mkdir(parents=True)
    path=folder/('source'+suffix);path.write_bytes(content)
    with Session() as db:
        c=ResearchContext(id=cid,name=Path(file.filename).stem[:200],summary=summary,source_text=text,source_name=Path(file.filename).name,source_path=str(path.resolve()))
        db.add(c);db.commit();return context_view(c)

@router.get('/contexts/{cid}/source')
def context_source(cid:str):
    with Session() as db:
        c=get(db,ResearchContext,cid)
        if not c.source_path:raise HTTPException(404,'Документ не загружен')
        return FileResponse(c.source_path,filename=c.source_name)

@router.get('/runs/{rid}/reports')
def reports(rid:str,company_id:str=''):
    with Session() as db:
        r=get(db,Run,rid)
        result={'context':r.context_snapshot or {'summary':r.sender_snapshot.get('profile',{}).get('description','')},
                'preparation':{'filename':r.filename,'rows':r.row_count,'companies':r.company_count,'sheets':r.config.get('sheets',[]),'missing_resources':db.query(Company).filter(Company.run_id==rid,Company.resource.like('unresolved:%')).count()},
                'segmentation':r.reports.get('segmentation'), 'segments':r.segments, 'stage':r.stage}
        if company_id:
            c=get(db,Company,company_id)
            if c.run_id!=rid:raise HTTPException(404,'Компания не относится к исследованию')
            result['company']=company_view(c,r,True)
            result['directions']={direction:{k:v for k,v in report.items() if k in ('discovery','sources','errors','state','summary','evidence','signals','contradictions','limitations')} for direction,report in c.research.items()}
        return result

@router.get('/runs/{rid}/segments')
def segment_summaries(rid:str):
    with Session() as db:
        r=get(db,Run,rid);result=[]
        for segment in r.segments:
            companies=db.query(Company).filter_by(run_id=rid,group=segment['id']).all()
            rows=db.query(SourceRow).join(Company,SourceRow.company_id==Company.id).filter(Company.run_id==rid,Company.group==segment['id']).count()
            strong=sum(sum(s.get('confidence')=='HIGH' and s.get('state')=='confirmed' for s in c.profile.get('signals',[])) for c in companies)
            conflicts=sum(bool(c.profile.get('contradictions')) for c in companies)
            result.append({**segment,'company_count':len(companies),'contact_count':rows,'strong_signals':strong,'conflicts':conflicts,'recommended':False})
        eligible=[s for s in result if s['company_count'] and s['contact_count'] and not s['conflicts']]
        if eligible:
            best=max(eligible,key=lambda s:(s['strong_signals']/s['company_count'],s.get('confidence')=='HIGH',s['contact_count']))
            best.update(recommended=True,recommendation_reason='Первый вариант для проверки: нет отмеченных противоречий; наиболее высокая средняя обеспеченность подтверждёнными признаками среди доступных сегментов. Должности и адреса получателей проверьте при выборе контактов.')
        return result

@router.get('/runs/{rid}/contacts')
def contacts(rid:str,segment_id:str='',company_id:str='',q:str='',offset:int=Query(0,ge=0),limit:int=Query(40,ge=1,le=100)):
    with Session() as db:
        r=get(db,Run,rid)
        query=db.query(SourceRow,Company).join(Company,SourceRow.company_id==Company.id).filter(SourceRow.run_id==rid)
        if segment_id:query=query.filter(Company.group==segment_id)
        if company_id:query=query.filter(Company.id==company_id)
        if q:query=query.filter(Company.name.ilike('%'+q+'%'))
        total=query.count();items=[]
        wb=load_workbook(r.source_path,read_only=False)
        try:
            for row,c in query.order_by(Company.name,SourceRow.sheet,SourceRow.row_index).offset(offset).limit(limit):
                ws=wb[row.sheet];fields=[{'column':cell.column,'name':str(ws.cell(1,cell.column).value or f'Колонка {cell.column}'),'value':str(cell.value) if cell.value is not None else ''} for cell in ws[row.row_index]]
                items.append({'id':row.id,'company_id':c.id,'company_name':c.name,'sheet':row.sheet,'row':row.row_index,'fields':fields})
        finally:wb.close()
        return {'items':items,'total':total}

class PackageInput(BaseModel):
    request_id:str=Field(default='',max_length=32,pattern=r'^([0-9a-f]{32})?$')
    name:str=Field(min_length=1,max_length=200)
    segment_id:str
    selection_mode:Literal['segment','companies','contacts']='segment'
    company_ids:list[str]=Field(default_factory=list,max_length=100000)
    row_ids:list[str]=Field(default_factory=list,max_length=100000)
    sender_id:str=''
    sender:SenderProfile|None=None
    save_sender:bool=False
    recipient_type:Literal['person','company']
    next_step:str=Field(min_length=1,max_length=1000)
    budget:float=0 # Legacy request field, no longer enforced.


def freeze_contacts(run,rows,path):
    if hashlib.sha256(Path(run.source_path).read_bytes()).hexdigest()!=run.source_hash:raise ValueError('Исходный файл изменён')
    source=load_workbook(run.source_path);out=Workbook();out.remove(out.active)
    try:
        for spec in run.config['sheets']:
            selected=[r for r in rows if r.sheet==spec['name']]
            if not selected:continue
            src=source[spec['name']];dst=out.create_sheet(spec['name']);dst.freeze_panes='A2'
            for key,value in src.column_dimensions.items():dst.column_dimensions[key]=copy(value)
            for target_i,source_i in enumerate([1]+[r.row_index for r in selected],1):
                for c in src[source_i]:
                    target=dst.cell(target_i,c.column,c.value);target.data_type=c.data_type
                    if c.has_style:target._style=copy(c._style)
                    if c.hyperlink:target._hyperlink=copy(c.hyperlink)
                if target_i>1 and fingerprint(dst[target_i])!=selected[target_i-2].fingerprint:raise ValueError('Изменена исходная строка')
            dst.auto_filter.ref=dst.dimensions
        out.save(path)
        check=load_workbook(path)
        try:
            for spec in run.config['sheets']:
                selected=[r for r in rows if r.sheet==spec['name']]
                if not selected:continue
                if fingerprint(check[spec['name']][1])!=fingerprint(source[spec['name']][1]):raise ValueError('Изменены заголовки')
                for i,row in enumerate(selected,2):
                    if fingerprint(check[row.sheet][i])!=row.fingerprint:raise ValueError('Изменена строка выгрузки')
        finally:check.close()
    finally:source.close();out.close()


def package_view(p,detail=False):
    result={k:getattr(p,k) for k in ('id','run_id','name','status','stage','error','budget','spent','reserved','created_at','recipient_type','next_step')}
    result.update(contact_count=len(p.selection['row_ids']),company_count=len(p.selection['company_ids']),segment_name=p.segment_snapshot['name'])
    if detail:result.update(segment=p.segment_snapshot,sender=p.sender_snapshot,context=p.context_snapshot,selection=p.selection,prompt=p.prompt,review=p.review)
    return result

@router.post('/runs/{rid}/packages')
def create_package(rid:str,payload:PackageInput):
    if not payload.name.strip() or not payload.next_step.strip():raise HTTPException(422,'Заполните название и следующий шаг')
    with Session() as db:
        r=get(db,Run,rid)
        request_hash=hashlib.sha256(json.dumps(payload.model_dump(exclude={'request_id'}),ensure_ascii=False,sort_keys=True).encode()).hexdigest()
        if payload.request_id and (existing:=db.get(PromptPackage,payload.request_id)):
            if existing.run_id!=rid or existing.selection.get('request_hash')!=request_hash:raise HTTPException(409,'Этот запрос уже создал другой комплект. Откройте список комплектов.')
            return package_view(existing)
        if r.status in ACTIVE or not r.segmentation_version:raise HTTPException(409,'Дождитесь завершения исследования и сегментации')
        if not credentials.for_run(r):raise HTTPException(409,'Добавьте API-ключ в настройках')
        segment=next((s for s in r.segments if s['id']==payload.segment_id),None)
        if not segment:raise HTTPException(404,'Сегмент не найден')
        query=db.query(SourceRow).join(Company,SourceRow.company_id==Company.id).filter(SourceRow.run_id==rid,Company.group==payload.segment_id)
        if payload.selection_mode=='companies':
            if not payload.company_ids:raise HTTPException(422,'Выберите компании')
            query=query.filter(SourceRow.company_id.in_(payload.company_ids))
        elif payload.selection_mode=='contacts':
            if not payload.row_ids or len(set(payload.row_ids))!=len(payload.row_ids):raise HTTPException(422,'Выберите контакты без повторного выбора одной строки')
            query=query.filter(SourceRow.id.in_(payload.row_ids))
        rows=query.order_by(SourceRow.sheet,SourceRow.row_index).all()
        if not rows:raise HTTPException(422,'В выборе нет контактов')
        if payload.selection_mode=='contacts' and {r.id for r in rows}!=set(payload.row_ids):raise HTTPException(422,'Контакты не относятся к выбранному сегменту')
        if payload.selection_mode=='companies' and {r.company_id for r in rows}!=set(payload.company_ids):raise HTTPException(422,'Компании не относятся к выбранному сегменту')
        if payload.sender_id:
            sender=get(db,Sender,payload.sender_id)
            if sender.archived:raise HTTPException(409,'Профиль удалён')
            snapshot={'id':sender.id,'version':sender.version,'profile':sender.profile}
        elif payload.sender:
            snapshot={'profile':payload.sender.model_dump()}
            if payload.save_sender:
                sender=Sender(profile=snapshot['profile']);db.add(sender);db.flush();snapshot.update(id=sender.id,version=sender.version)
        else:raise HTTPException(422,'Выберите или заполните отправителя')
        if not snapshot['profile']['brand'].strip():raise HTTPException(422,'Укажите бренд отправителя')
        pid=payload.request_id or uid();folder=(settings.data_dir/'packages'/pid).resolve();folder.mkdir(parents=True)
        try:
            freeze_contacts(r,rows,folder/'contacts.xlsx')
            selection={'row_ids':[row.id for row in rows],'company_ids':sorted({row.company_id for row in rows}),'source_hash':r.source_hash,'segmentation_version':r.segmentation_version,'contacts_sha256':hashlib.sha256((folder/'contacts.xlsx').read_bytes()).hexdigest(),'request_hash':request_hash}
            p=PromptPackage(id=pid,run_id=rid,name=payload.name.strip(),segment_snapshot={k:v for k,v in segment.items() if k!='prompt'},sender_snapshot=snapshot,context_snapshot=r.context_snapshot or {'summary':r.sender_snapshot.get('profile',{}).get('description','')},selection=selection,recipient_type=payload.recipient_type,next_step=payload.next_step.strip(),budget=0,path=str(folder))
            db.add(p);db.commit();return package_view(p)
        except ValueError as e:
            db.rollback();shutil.rmtree(folder,ignore_errors=True);raise HTTPException(409,str(e))
        except Exception:
            db.rollback();shutil.rmtree(folder,ignore_errors=True);raise

@router.get('/runs/{rid}/packages')
def packages(rid:str):
    with Session() as db:
        get(db,Run,rid)
        return [package_view(p) for p in db.query(PromptPackage).filter_by(run_id=rid).order_by(PromptPackage.created_at.desc())]

@router.get('/packages/{pid}')
def package_detail(pid:str):
    with Session() as db:return package_view(get(db,PromptPackage,pid),True)

@router.post('/packages/{pid}/pause')
def pause_package(pid:str):
    with Session() as db:
        p=get(db,PromptPackage,pid)
        if p.status not in ('QUEUED','RUNNING'):raise HTTPException(409,'Комплект не обрабатывается')
        p.status='PAUSED' if p.status=='QUEUED' else 'PAUSING';db.commit();return package_view(p)

@router.post('/packages/{pid}/resume')
def resume_package(pid:str):
    with Session() as db:
        p=get(db,PromptPackage,pid)
        if p.status not in ('PAUSED','FAILED'):raise HTTPException(409,'Комплект нельзя продолжить в текущем состоянии')
        if not credentials.for_run(None):raise HTTPException(409,'Добавьте API-ключ в настройках')
        p.status='QUEUED';p.error=None;db.commit();return package_view(p)

@router.get('/packages/{pid}/download')
def package_download(pid:str,file:Literal['contacts.xlsx','prompt.md','package.zip']='package.zip'):
    with Session() as db:
        p=get(db,PromptPackage,pid)
        if p.status!='COMPLETED':raise HTTPException(409,'Комплект ещё не готов')
        return FileResponse(Path(p.path)/file,filename=file)


def generate_package(pid,gateway=None):
    with Session() as db:
        p=db.get(PromptPackage,pid);rid=p.run_id
        if p.status!='RUNNING':raise RunPaused()
        sender={k:v for k,v in p.sender_snapshot['profile'].items() if k not in ('offers','description','materials')}
        payload={'segment':p.segment_snapshot,'sender':sender,'research_context':p.context_snapshot.get('summary',''),'recipient_type':p.recipient_type,'next_step':p.next_step,'contacts':len(p.selection['row_ids'])}
        prompt=p.prompt;review=p.review
    gw=gateway or Gateway(rid,package_id=pid)
    if not prompt:
        result=gw.ask('prompt',PROMPT+' Тип person означает обращение к сотруднику, company — к ответственному представителю компании без выдуманного имени. Используй выбранный следующий шаг. Если данные отправителя пустые, не придумывай их. Предложение формулируй на основе research_context и гипотезы сегмента.',payload,PromptResult,max_output=7000)
        prompt=result.text
        if not all(x in prompt for x in ('INPUT_ROW_DATA','Письмо 1','Письмо 2','Письмо 3','Письмо 4')):raise ValueError('Промпт не прошёл структурную проверку')
        representative=payload['sender'].get('representative','').strip()
        if representative and representative not in prompt:raise ValueError('В промпте не сохранено имя отправителя')
        with Session() as db:
            p=db.get(PromptPackage,pid);p.prompt=prompt;p.stage='REVIEW';db.commit()
    if not review.get('passed'):
        checked=gw.ask('synthesis',REVIEW,{**payload,'prompt':prompt},Review,max_output=3000)
        with Session() as db:
            p=db.get(PromptPackage,pid);p.review=checked.model_dump();db.commit()
        if not checked.passed:raise ValueError('Промпт не принят проверкой: '+'; '.join(checked.issues))
    with Session() as db:
        p=db.get(PromptPackage,pid);folder=Path(p.path)
        if p.status!='RUNNING':raise RunPaused()
        if hashlib.sha256((folder/'contacts.xlsx').read_bytes()).hexdigest()!=p.selection['contacts_sha256']:raise ValueError('Файл выбранных контактов изменён')
        (folder/'prompt.md').write_text(prompt,encoding='utf-8')
        (folder/'snapshot.json').write_text(json.dumps({'segment':p.segment_snapshot,'sender':p.sender_snapshot,'context':p.context_snapshot,'selection':p.selection,'recipient_type':p.recipient_type,'next_step':p.next_step,'review':p.review},ensure_ascii=False,indent=2),encoding='utf-8')
        with zipfile.ZipFile(folder/'package.zip','w',zipfile.ZIP_DEFLATED) as z:
            for name in ('contacts.xlsx','prompt.md','snapshot.json'):z.write(folder/name,name)
        p.status='COMPLETED';p.stage='READY';p.error=None
        log(db,rid,'PROMPT_PACKAGE','Подготовлен отдельный комплект для генерации',metrics={'package_id':pid,'contacts':len(p.selection['row_ids'])});db.commit()
