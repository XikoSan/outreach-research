from .audiences import GENERAL, GENERAL_RULES, audience_filter, general_segment, general_template
"""Immutable workbook registry and verified row-preserving partition export."""
import hashlib, json, re, shutil, unicodedata, zipfile
from copy import copy
from pathlib import Path
from urllib.parse import urlsplit, urlunsplit
from collections import defaultdict
from datetime import date, datetime, time
from openpyxl import load_workbook, Workbook
from .db import Company, SourceRow, Export, uid

EMPTY = {'', '-', '—', 'нет', 'null', 'none', 'n/a', 'nan'}
SOCIAL = {'vk.com','vkontakte.ru','linkedin.com','facebook.com','instagram.com','t.me','telegram.me','ok.ru','youtube.com','x.com','twitter.com','tiktok.com','pinterest.com','threads.net','dzen.ru','zen.yandex.ru','rutube.ru'}
SOCIAL_LABELS = ('соцсет', 'соц сет', 'социальн', 'профил', 'страниц', 'social', 'profile', 'page', 'vk', 'вк', 'vkontakte', 'linkedin', 'линкедин', 'telegram', 'телеграм', 'instagram', 'инстаграм', 'facebook', 'фейсбук', 'однокласс', 'youtube', 'ютуб', 'twitter', 'твиттер', 'x.com', 'tiktok', 'тикток', 'pinterest', 'пинтерест', 'threads', 'dzen', 'дзен', 'rutube', 'рутуб')
COMPANY_LABELS = ('компан', 'организац', 'бренд', 'официальн', 'company', 'corporate', 'business', 'official')
PERSON_LABELS = ('сотрудник', 'контактного лица', 'личн', 'персональн', 'менеджер', 'директор', 'руководител', 'employee', 'contact person', 'personal', 'manager', 'owner', 'founder', 'ceo')
SHARED = {'yandex.ru','ya.ru','2gis.ru','google.com','maps.google.com','tilda.ws','wixsite.com'}

def clean(v):
    s = str(v or '').strip()
    return '' if s.lower() in EMPTY else s

def canonical_resource(value, social=False):
    value=clean(value)
    if not value: return None
    value=re.split(r'[\s;,]+',value)[0]
    try:
        p=urlsplit(value if '://' in value else 'https://'+value)
        host=(p.hostname or '').lower().removeprefix('www.').removeprefix('m.').encode('idna').decode()
        if p.scheme not in ('http','https') or '.' not in host or p.username or p.password or p.port not in (None,80,443): return None
        if host in SHARED or (social and host=='youtu.be'): return None
        is_social=host in SOCIAL or bool(social and p.path.strip('/'))
        if is_social:
            parts=p.path.strip('/').split('/')
            if not parts[0]:return None
            # Posts, videos and personal LinkedIn profiles are not company pages.
            excluded={
                'linkedin.com':{'in','posts','feed'},
                'youtube.com':{'watch','shorts','playlist','live'},
                'instagram.com':{'p','reel','stories'},
                'facebook.com':{'watch','photo','posts','reel'},
                'vk.com':{'wall','feed'},
                't.me':{'c','joinchat'},
                'telegram.me':{'c','joinchat'},
                'pinterest.com':{'pin'},
                'x.com':{'home','i','search'},
                'twitter.com':{'home','i','search'},
            }
            if parts[0].casefold() in excluded.get(host,set()):return None
            if host in {'vk.com','vkontakte.ru'} and parts[0].casefold().startswith('wall'):return None
            if host in {'x.com','twitter.com','tiktok.com'} and len(parts)>1:return None
            if host in {'t.me','telegram.me'} and parts[0].startswith('+'):return None
        # Retain full tenant paths on corporate social resources.
        path=p.path.rstrip('/') if is_social else ''
        if social and not is_social:return None
        return urlunsplit(('https',host,path,'',''))
    except (ValueError,UnicodeError):return None

def corporate_social_columns(headers):
    """Only company-labelled social fields may supply a provisional identity anchor."""
    return [i for i, header in enumerate(headers)
            if (label := header.casefold().replace('ё', 'е'))
            and (any(token in label for token in SOCIAL_LABELS) or bool(re.search(r'(?<![a-zа-я0-9])x(?![a-zа-я0-9])',label)))
            and any(token in label for token in COMPANY_LABELS)
            and not any(token in label for token in PERSON_LABELS)]

def corporate_socials(cells, columns):
    links=[]
    for index in columns:
        cell=cells[index]
        value=clean(cell.hyperlink.target if cell.hyperlink else cell.value)
        for part in re.split(r'[\s;,|]+',value):
            url=canonical_resource(part,True)
            if url and url not in links:links.append(url)
    return links

def fingerprint(cells):
    vals=[]
    for c in cells:
        v=c.value
        if v is None or v=='':
            vals.append(['empty','str',''])
            continue
        if isinstance(v,(datetime,date,time)):v=v.isoformat()
        vals.append([c.data_type,type(c.value).__name__,v])
    return hashlib.sha256(json.dumps(vals,ensure_ascii=False,default=str).encode()).hexdigest()

def preflight(path):
    try:
        with zipfile.ZipFile(path) as z:
            if sum(x.file_size for x in z.infolist())>200_000_000: raise ValueError('Распакованный XLSX превышает 200 МБ')
            if any('externalLink' in x.filename for x in z.infolist()):raise ValueError('Внешние связи Excel не поддерживаются')
        wb=load_workbook(path)
    except ValueError:raise
    except Exception as e:raise ValueError('Не удалось открыть XLSX') from e
    data=[]
    for ws in wb:
        if ws.max_row>100_001 or ws.max_column>200: raise ValueError('Лимит: 100 000 строк и 200 колонок на лист')
        if ws.merged_cells.ranges: raise ValueError(f'{ws.title}: объединённые ячейки не поддерживаются')
        if ws.max_row<2:continue
        headers=[clean(c.value) for c in ws[1]]
        names={h.casefold():i for i,h in enumerate(headers) if h}
        company=next((names[h] for h in ['название компании','company','company name'] if h in names),None)
        if company is None:raise ValueError(f'{ws.title}: первая строка должна содержать «Название компании»')
        rows=[]
        for cells in ws.iter_rows(min_row=2):
            if any(c.data_type=='f' for c in cells):raise ValueError(f'{ws.title}: формулы не поддерживаются, сохраните вычисленные значения')
            if any(c.value is not None for c in cells):rows.append(cells)
        if rows:data.append((ws,headers,names,company,rows))
    if not data:raise ValueError('В файле нет строк контактов')
    return wb,data

def company_name_key(value):
    """Use the displayed company name, ignoring case, punctuation and legal-form prefix."""
    value=unicodedata.normalize('NFKC',clean(value)).casefold().replace('ё','е')
    value=re.sub(r'^(?:(?:ооо|оао|зао|пао|ао|ип|llc|ltd)\s+)+','',value)
    return ' '.join(re.findall(r'[a-zа-я0-9]+',value))

def normalize_inn(value):
    value=clean(value).replace('\xa0','').strip()
    if re.fullmatch(r'\d+\.0',value):value=value[:-2]
    value=re.sub(r'[\s-]+','',value)
    return value if value.isdigit() else value.casefold()

def import_rows(db,run):
    wb,sheets=preflight(run.source_path)
    entries=[]
    for ws,headers,names,name_col,rows in sheets:
        social_columns=corporate_social_columns(headers)
        def val(cells,*keys):
            for key in keys:
                if key.casefold() in names:
                    cell=cells[names[key.casefold()]]
                    value=clean(cell.hyperlink.target if cell.hyperlink else cell.value)
                    if value:return value
            return ''
        for cells in rows:
            name=clean(cells[name_col].value) or 'Компания не указана'
            website=canonical_resource(val(cells,'Ссылка на сайт','Website','Сайт'))
            socials=corporate_socials(cells,social_columns)
            if website and canonical_resource(website,True):
                socials=list(dict.fromkeys([website,*socials]));website=None
            name_key=company_name_key(name)
            if not name_key or name=='Компания не указана':name_key=f'unknown:{ws.title}:{cells[0].row}'
            entries.append({'name':name,'name_key':name_key,
                'inn':normalize_inn(val(cells,'ИНН','Tax ID','ИНН компании')),
                'legal_name':val(cells,'Юридическое название компании','Legal company name','Legal name'),
                'website':website,'socials':socials,'description':val(cells,'Описание'),
                'sheet':ws.title,'row_index':cells[0].row,'cells':cells,'headers':headers})
    groups=[];by_name=defaultdict(list);by_name_inn={}
    def new_group(entry):
        group={'name_key':entry['name_key'],'inn':entry['inn'],'entries':[],'ambiguous':False}
        groups.append(group);by_name[entry['name_key']].append(group)
        return group
    # Resolve rows with an INN first, so blank-INN rows can be assigned independent of row order.
    for entry in entries:
        if not entry['inn']:continue
        key=(entry['name_key'],entry['inn'])
        group=by_name_inn.get(key)
        if group is None:
            group=new_group(entry);by_name_inn[key]=group
        group['entries'].append(entry)
    blank_groups={}
    for entry in entries:
        if entry['inn']:continue
        candidates=[group for group in by_name[entry['name_key']] if group['inn']]
        group=None
        if len(candidates)==1:
            group=candidates[0]
            known_sites={item['website'] for item in group['entries'] if item['website']}
            if entry['website'] and known_sites and entry['website'] not in known_sites:
                group['ambiguous']=True
        elif len(candidates)>1:
            links=set(([entry['website']] if entry['website'] else [])+entry['socials'])
            matches=[candidate for candidate in candidates if links & {
                link for item in candidate['entries'] for link in
                (([item['website']] if item['website'] else [])+item['socials'])}]
            if len(matches)==1:group=matches[0]
        if group is None:
            group=blank_groups.get(entry['name_key'])
            if group is None:
                group=new_group(entry);blank_groups[entry['name_key']]=group
            if len(candidates)>1:group['ambiguous']=True
        group['entries'].append(entry)
    prepared=[];by_resource=defaultdict(list)
    for group in groups:
        members=group['entries']
        websites=list(dict.fromkeys(item['website'] for item in members if item['website']))
        socials=list(dict.fromkeys(link for item in members for link in item['socials']))
        resource=(websites or socials or ['unresolved:'+uid()])[0]
        if not group['inn'] and len(websites)>1:group['ambiguous']=True
        prepared.append((group,websites,socials,resource))
        if not resource.startswith('unresolved:'):by_resource[resource].append(group)
    companies=[];rows_out=[]
    for group,websites,socials,resource in prepared:
        members=group['entries']
        shared=len(by_resource.get(resource,[]))>1
        conflict=group['ambiguous'] or shared
        names=list(dict.fromkeys(item['name'] for item in members))
        legal_names=list(dict.fromkeys(item['legal_name'] for item in members if item['legal_name']))
        legal_ids=list(dict.fromkeys(item['inn'] for item in members if item['inn']))
        description=next((item['description'] for item in members if item['description']),'')
        reason=('Один ресурс указан у разных компаний' if shared else
                'Неоднозначно, к какой компании относится часть контактов' if conflict else
                'Нет подтверждаемого сайта или корпоративной соцстраницы' if resource.startswith('unresolved:') else '')
        needs_review=conflict or resource.startswith('unresolved:')
        company=Company(id=uid(),run_id=run.id,name=names[0],resource=resource,
            context={'names':names,'legal_names':legal_names,'legal_ids':legal_ids,
                     'websites':websites,'socials':socials,'description':description,
                     'identity_conflict':conflict,'grouping_basis':'company_name_inn'},
            identity='NEEDS_REVIEW' if needs_review else 'UNVERIFIED',
            group='NEEDS_REVIEW' if needs_review else 'NOT_PROCESSED',reason=reason)
        companies.append(company)
        for entry in members:
            cells=entry['cells'];headers=entry['headers']
            rows_out.append(SourceRow(id=uid(),run_id=run.id,company_id=company.id,
                sheet=entry['sheet'],row_index=entry['row_index'],fingerprint=fingerprint(cells),
                display_fields=[{'column':cell.column,
                    'name':str(headers[cell.column-1] or f'Колонка {cell.column}'),
                    'value':str(cell.value) if cell.value is not None else ''} for cell in cells]))
    db.add_all(companies);db.flush();db.add_all(rows_out)
    from .base_stats import index_companies
    index_companies(db,run,companies)
    run.row_count=len(rows_out);run.company_count=len(companies)
    run.config={**run.config,'sheets':[{'name':w.title,'headers':h,'rows':len(rs)} for w,h,_,_,rs in sheets]}
    wb.close()

STATUS_GROUPS={'EXCLUDED','FAILED','NEEDS_REVIEW','INSUFFICIENT_DATA','AMBIGUOUS','UNCLASSIFIED','NOT_PROCESSED'}

def export_run(db,run,root,partial=None):
    if hashlib.sha256(Path(run.source_path).read_bytes()).hexdigest()!=run.source_hash:raise ValueError('Исходный файл изменён')
    companies={c.id:c for c in db.query(Company).filter_by(run_id=run.id)}
    rows=db.query(SourceRow).filter_by(run_id=run.id).order_by(SourceRow.sheet,SourceRow.row_index).all()
    segments={s['id']:s for s in run.segments}
    grouped=defaultdict(list)
    for r in rows:
        group=companies[r.company_id].group
        if group not in STATUS_GROUPS and not re.fullmatch(r'segment_\d{3,}',group):raise ValueError('Некорректная группа')
        grouped[group if group in segments else GENERAL].append(r)
    export_id=uid();stage=Path(root)/('tmp-'+export_id);dest=Path(root)/export_id
    stage.mkdir(parents=True)
    source=load_workbook(run.source_path)
    manifest={'rows':len(rows),'files':[],'groups':{},'segmentation_version':run.segmentation_version,'partial':(run.status not in ('COMPLETED','COMPLETED_WITH_ISSUES')) if partial is None else partial}
    seen=set()
    try:
        for group,members in sorted(grouped.items()):
            folder='databases'
            out=Workbook();out.remove(out.active)
            bysheet=defaultdict(list)
            for row in members:bysheet[row.sheet].append(row)
            for spec in run.config['sheets']:
                title=spec['name'];src=source[title];ws=out.create_sheet(title)
                ws.freeze_panes='A2'
                for key,dim in src.column_dimensions.items():ws.column_dimensions[key]=copy(dim)
                indices=[1]+[r.row_index for r in bysheet[title]]
                for dest_i,original_i in enumerate(indices,1):
                    for c in src[original_i]:
                        target=ws.cell(dest_i,c.column,c.value)
                        target.data_type=c.data_type
                        if c.has_style:target._style=copy(c._style)
                        if c.hyperlink:target._hyperlink=copy(c.hyperlink)
                        if c.comment:target.comment=copy(c.comment)
                    ws.row_dimensions[dest_i].height=src.row_dimensions[original_i].height
                ws.auto_filter.ref=ws.dimensions
            if group == GENERAL:
                audit=out.create_sheet('Причины группы Общее')
                audit.append(['Компания', 'Причина', 'Основание', 'Ссылка'])
                unique={row.company_id for row in members}
                for company_id in sorted(unique):
                    company=companies[company_id]
                    exclusion=company.profile.get('exclusion', {})
                    urls=exclusion.get('urls', [])
                    audit.append([company.name, exclusion.get('reason', company.reason),
                                  exclusion.get('note', ''), urls[0] if urls else ''])
                    for cell in audit[audit.max_row]:
                        cell.data_type='s'
                audit.freeze_panes='A2'
                audit.auto_filter.ref=audit.dimensions
                for column,width in {'A':36,'B':70,'C':70,'D':70}.items():
                    audit.column_dimensions[column].width=width
            target=stage/folder/(group+'.xlsx');target.parent.mkdir(exist_ok=True);out.save(target);out.close()
            verified=load_workbook(target)
            for spec in run.config['sheets']:
                title=spec['name'];check=verified[title]
                if fingerprint(check[1])!=fingerprint(source[title][1]):raise ValueError('Изменены заголовки')
                for idx,row in enumerate(bysheet[title],2):
                    if fingerprint(check[idx])!=row.fingerprint:raise ValueError(f'Изменена строка {row.id}')
                    if row.id in seen:raise ValueError('Повтор строки в экспорте')
                    seen.add(row.id)
            verified.close();manifest['files'].append(str(target.relative_to(stage)));manifest['groups'][group]=len(members)
            if group == GENERAL:
                p=stage/'prompts'/'GENERAL.md';p.parent.mkdir(exist_ok=True);p.write_text(general_template(run),encoding='utf-8');manifest['files'].append(str(p.relative_to(stage)))
            if group not in STATUS_GROUPS and segments.get(group,{}).get('prompt'):
                p=stage/'prompts'/(group+'.md');p.parent.mkdir(exist_ok=True);p.write_text(segments[group]['prompt'],encoding='utf-8');manifest['files'].append(str(p.relative_to(stage)))
        if len(seen)!=run.row_count:raise ValueError('Нарушено количество строк')
        manifest['checksums']={f:hashlib.sha256((stage/f).read_bytes()).hexdigest() for f in manifest['files']}
        with zipfile.ZipFile(stage/'results.zip','w',zipfile.ZIP_DEFLATED) as z:
            for f in manifest['files']:z.write(stage/f,f)
        stage.rename(dest)
        record=Export(id=export_id,run_id=run.id,path=str(dest),manifest=manifest);db.add(record);db.commit();return record
    except Exception:
        shutil.rmtree(stage,ignore_errors=True);raise
    finally:source.close()
