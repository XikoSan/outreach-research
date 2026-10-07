from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import re

ROOT = Path('/home/hiko/Project/outreach-research')
doc = Document()
sec = doc.sections[0]
sec.page_height, sec.page_width = Cm(29.7), Cm(21)
sec.top_margin, sec.bottom_margin = Cm(1.7), Cm(1.65)
sec.left_margin, sec.right_margin = Cm(1.9), Cm(1.9)
sec.header_distance, sec.footer_distance = Cm(.7), Cm(.7)
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'List Bullet', 'List Number']:
    s = doc.styles[name]
    s.font.name = 'DejaVu Sans'
    s.font.color.rgb = RGBColor(0,0,0)
    s.paragraph_format.widow_control = True
normal = doc.styles['Normal']
normal.font.size = Pt(10)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.08
for name, size in [('Title',27),('Heading 1',18),('Heading 2',11.5)]:
    s=doc.styles[name]
    s.font.size=Pt(size)
    s.font.bold=True
    s.paragraph_format.space_before=Pt(10 if name=='Heading 2' else 0)
    s.paragraph_format.space_after=Pt(7)
    s.paragraph_format.keep_with_next=True
doc.styles['Subtitle'].font.size=Pt(15)
for name in ['List Bullet','List Number']:
    doc.styles[name].paragraph_format.space_after=Pt(4)
    doc.styles[name].paragraph_format.line_spacing=1.04

footer=sec.footer.paragraphs[0]
footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=footer.add_run('Research and Segmentation  •  ')
r.font.size=Pt(8)
fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); footer._p.append(fld)

def table(rows):
    t=doc.add_table(rows=1, cols=len(rows[0]))
    t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.autofit=False
    widths=([5.1,12.1] if len(rows[0])==2 else [5.2,7.1,4.9])
    for i,w in enumerate(widths): t.columns[i].width=Cm(w)
    pr=t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+edge); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),'D9D9D9'); borders.append(e)
    pr.append(borders)
    for ri,row in enumerate(rows):
        cells=t.rows[0].cells if ri==0 else t.add_row().cells
        trpr=t.rows[ri]._tr.get_or_add_trPr()
        cant=OxmlElement('w:cantSplit'); trpr.append(cant)
        if ri==0:
            repeat=OxmlElement('w:tblHeader'); trpr.append(repeat)
        for ci,txt in enumerate(row):
            c=cells[ci]; c.width=Cm(widths[ci]); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp=c._tc.get_or_add_tcPr()
            mar=OxmlElement('w:tcMar')
            for side in ['top','bottom','left','right']:
                e=OxmlElement('w:'+side); e.set(qn('w:w'),'90'); e.set(qn('w:type'),'dxa'); mar.append(e)
            cp.append(mar)
            shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'23364A' if ri==0 else ('F1F4F7' if ri%2==0 else 'FFFFFF')); cp.append(shade)
            p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.02
            r=p.add_run(txt); r.font.size=Pt(8.5); r.bold=(ri==0)
            if ri==0:r.font.color.rgb=RGBColor(255,255,255)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)

lines=(ROOT/'development_spec_source.md').read_text().splitlines()
i=0; first=True
while i<len(lines):
    line=lines[i].strip()
    if not line: i+=1; continue
    if line.startswith('|'):
        rows=[]
        while i<len(lines) and lines[i].strip().startswith('|'):
            rows.append([s.strip() for s in lines[i].strip().strip('|').split('|')]); i+=1
        table(rows); continue
    if line.startswith('# '):
        if first:
            doc.add_paragraph(line[2:],'Title');first=False
        else:
            p=doc.add_paragraph(line[2:],'Heading 1');p.paragraph_format.page_break_before=True
    elif line.startswith('## '): doc.add_paragraph(line[3:],'Heading 2')
    elif line.startswith('- '):doc.add_paragraph(line[2:],'List Bullet')
    elif re.match(r'^\d+\. ',line):doc.add_paragraph(line,'Normal')
    elif line=='Research and Segmentation Tool для B2B outbound':doc.add_paragraph(line,'Subtitle')
    else:doc.add_paragraph(line)
    i+=1
doc.core_properties.title='Техническое задание Research and Segmentation Tool'
doc.core_properties.subject='Единая спецификация разработки версии 2.0'
doc.core_properties.author=''
out=ROOT/'deliverables'/'TZ_Research_Segmentation_v2.docx'
out.parent.mkdir(exist_ok=True)
for el in doc.styles.element.xpath('.//w:pBdr'):
    el.getparent().remove(el)
for el in doc.element.xpath('.//w:pBdr'):
    el.getparent().remove(el)
doc.save(out)
print(out)
